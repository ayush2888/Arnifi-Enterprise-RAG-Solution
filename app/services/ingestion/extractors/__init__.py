"""Extract plain text (and row texts) from local Drive artifacts."""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pypdf import PdfReader

from app.utils.helpers import clean_text, get_logger

logger = get_logger(__name__)

PDF_MIME = "application/pdf"
CSV_MIME = "text/csv"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

SUPPORTED_MIME_TYPES = frozenset({PDF_MIME, CSV_MIME, XLSX_MIME})


def is_supported_drive_file(mime_type: str, name: str = "") -> bool:
    mime = (mime_type or "").lower()
    lower_name = (name or "").lower()
    if mime in SUPPORTED_MIME_TYPES:
        return True
    return (
        lower_name.endswith(".pdf")
        or lower_name.endswith(".csv")
        or lower_name.endswith(".xlsx")
    )


def resolve_extract_kind(mime_type: str, name: str = "") -> str | None:
    mime = (mime_type or "").lower()
    lower_name = (name or "").lower()
    if mime == PDF_MIME or lower_name.endswith(".pdf"):
        return "pdf"
    if mime == CSV_MIME or lower_name.endswith(".csv"):
        return "csv"
    if mime == XLSX_MIME or lower_name.endswith(".xlsx"):
        return "xlsx"
    return None


def extract_pdf_text(path: Path) -> str:
    reader = PdfReader(str(path))
    parts: list[str] = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        if page_text.strip():
            parts.append(page_text)
    return clean_text("\n\n".join(parts))


def _decode_bytes(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


# Pathological sheets (e.g. 25k penalty rows) can stall Titan embed for hours.
_MAX_TABULAR_ROWS = 5000


def _row_texts_from_matrix(
    matrix: list[list[str]],
    *,
    sheet_name: str | None = None,
    max_rows: int = _MAX_TABULAR_ROWS,
) -> list[str]:
    """Turn a 2D cell matrix into one embeddable string per data row."""
    if not matrix:
        return []

    headers_raw = matrix[0]
    headers: list[str] = []
    for i, h in enumerate(headers_raw):
        label = (h or "").strip()
        headers.append(label if label else f"col_{i + 1}")

    out: list[str] = []
    skipped = 0
    for row in matrix[1:]:
        if not any((c or "").strip() for c in row):
            continue
        if len(out) >= max_rows:
            skipped += 1
            continue
        pairs: list[str] = []
        for i, header in enumerate(headers):
            val = (row[i] if i < len(row) else "").strip()
            if val:
                pairs.append(f"{header}: {val}")
        if not pairs:
            continue
        body = "\n".join(pairs)
        if sheet_name:
            out.append(f"Sheet: {sheet_name}\n{body}")
        else:
            out.append(body)
    if skipped:
        logger.warning(
            "Tabular extract capped at %d rows (skipped %d more)%s",
            max_rows,
            skipped,
            f" on sheet {sheet_name!r}" if sheet_name else "",
        )
    return out


def extract_csv_rows(path: Path) -> list[str]:
    text = _decode_bytes(path.read_bytes())
    reader = csv.reader(text.splitlines())
    matrix = [[(cell or "").strip() for cell in row] for row in reader]
    # Drop fully empty trailing lines already handled; keep sparse rows with some cells.
    return _row_texts_from_matrix(matrix)


def extract_csv_text(path: Path) -> str:
    """Flat preview text (also used when row list is joined)."""
    rows = extract_csv_rows(path)
    if rows:
        return clean_text("\n\n".join(rows))
    # Fallback for header-only / malformed files: pipe join like before.
    raw = _decode_bytes(path.read_bytes())
    lines: list[str] = []
    for row in csv.reader(raw.splitlines()):
        cells = [cell.strip() for cell in row if cell and cell.strip()]
        if cells:
            lines.append(" | ".join(cells))
    return clean_text("\n".join(lines))


def extract_xlsx_rows(path: Path) -> list[str]:
    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    all_rows: list[str] = []
    try:
        for sheet in workbook.worksheets:
            matrix: list[list[str]] = []
            for excel_row in sheet.iter_rows(values_only=True):
                matrix.append([_cell_str(cell) for cell in excel_row])
            all_rows.extend(_row_texts_from_matrix(matrix, sheet_name=sheet.title))
    finally:
        workbook.close()
    return all_rows


def extract_xlsx_text(path: Path) -> str:
    rows = extract_xlsx_rows(path)
    return clean_text("\n\n".join(rows))


def extract_drive_content(
    path: Path,
    mime_type: str = "",
    name: str = "",
) -> dict[str, Any]:
    """
    Extract Drive file content for sync artifacts.

    Returns kind, flat text (preview / PDF body), and optional row_texts for
    tabular sources (one string per data row).
    """
    kind = resolve_extract_kind(mime_type, name or path.name)
    if kind == "pdf":
        text = extract_pdf_text(path)
        return {"kind": kind, "text": text, "row_texts": None}
    if kind == "csv":
        row_texts = extract_csv_rows(path)
        text = clean_text("\n\n".join(row_texts)) if row_texts else extract_csv_text(path)
        return {"kind": kind, "text": text, "row_texts": row_texts or None}
    if kind == "xlsx":
        row_texts = extract_xlsx_rows(path)
        text = clean_text("\n\n".join(row_texts))
        return {"kind": kind, "text": text, "row_texts": row_texts or None}
    raise ValueError(f"Unsupported file for extraction: mime={mime_type!r} name={name!r}")


def extract_text(path: Path, mime_type: str = "", name: str = "") -> str:
    return str(extract_drive_content(path, mime_type=mime_type, name=name)["text"])
